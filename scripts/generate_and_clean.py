"""
Sales Data Analysis - data generation, cleaning and Excel report builder.
NOTE: The dataset is SYNTHETIC (randomly generated with a fixed seed) for portfolio/demo purposes.
Run from the project root:  python scripts/generate_and_clean.py
"""
import numpy as np, pandas as pd, os, pickle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils import get_column_letter

np.random.seed(42)
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA, XL = os.path.join(ROOT, "data"), os.path.join(ROOT, "excel")

# ---------- 1. Generate raw (messy) data ----------
products = {  # product: (category, price, cost, popularity)
 "Laptop":("Electronics",55000,42000,5),"Smartphone":("Electronics",28000,21000,8),
 "Headphones":("Electronics",3500,2100,10),"Smartwatch":("Electronics",9000,6000,6),
 "Office Chair":("Furniture",8500,5500,5),"Study Desk":("Furniture",12000,8000,4),
 "Bookshelf":("Furniture",6500,4200,3),"Notebook Pack":("Stationery",450,250,14),
 "Pen Set":("Stationery",300,140,15),"Backpack":("Accessories",1800,950,9),
 "Water Bottle":("Accessories",600,280,12),"Desk Lamp":("Accessories",1400,750,7)}
regions = {"North":0.22,"South":0.27,"East":0.16,"West":0.25,"Central":0.10}
segments = ["Consumer","Corporate","Small Business"]
n = 5000
dates = pd.date_range("2025-01-01","2026-05-31")
season = {1:.9,2:.85,3:1.0,4:1.0,5:1.05,6:.95,7:.9,8:1.0,9:1.05,10:1.3,11:1.4,12:1.25}
w = np.array([season[d.month]*(1+0.02*((d.year-2025)*12+d.month)) for d in dates]); w/=w.sum()
pn = list(products); pw = np.array([products[p][3] for p in pn],float); pw/=pw.sum()
df = pd.DataFrame({
 "Order_ID":[f"ORD-{100000+i}" for i in range(n)],
 "Order_Date":np.random.choice(dates,n,p=w),
 "Region":np.random.choice(list(regions),n,p=list(regions.values())),
 "Product":np.random.choice(pn,n,p=pw),
 "Customer_Segment":np.random.choice(segments,n,p=[.55,.3,.15]),
 "Quantity":np.random.choice([1,2,3,4,5],n,p=[.5,.25,.12,.08,.05]),
 "Discount":np.random.choice([0,0.05,0.10,0.15,0.20],n,p=[.45,.2,.2,.1,.05])})
df["Category"]=df.Product.map(lambda p:products[p][0])
df["Unit_Price"]=df.Product.map(lambda p:products[p][1])
df["Unit_Cost"]=df.Product.map(lambda p:products[p][2])
raw=df[["Order_ID","Order_Date","Region","Category","Product","Customer_Segment","Quantity","Unit_Price","Unit_Cost","Discount"]].copy()
raw["Order_Date"]=pd.to_datetime(raw["Order_Date"])
# inject data-quality problems
idx=np.random.choice(n,120,replace=False)
raw.loc[idx[:40],"Region"]=raw.loc[idx[:40],"Region"].str.lower()
raw.loc[idx[40:80],"Region"]=" "+raw.loc[idx[40:80],"Region"]+" "
raw.loc[idx[80:100],"Region"]=np.nan
raw.loc[idx[100:120],"Discount"]=np.nan
raw["Order_Date"]=raw["Order_Date"].dt.strftime("%Y-%m-%d")
alt=np.random.choice(n,150,replace=False)
raw.loc[alt,"Order_Date"]=pd.to_datetime(raw.loc[alt,"Order_Date"]).dt.strftime("%d/%m/%Y")
raw=pd.concat([raw,raw.sample(60,random_state=1)],ignore_index=True)   # duplicates
raw.to_csv(os.path.join(DATA,"raw_sales_data.csv"),index=False)

# ---------- 2. Clean ----------
log=[f"Raw rows: {len(raw)}"]
c=pd.read_csv(os.path.join(DATA,"raw_sales_data.csv"))
d0=len(c); c=c.drop_duplicates(); log.append(f"Duplicate rows removed: {d0-len(c)}")
c["Order_Date"]=c["Order_Date"].apply(lambda s: pd.to_datetime(s,format="%Y-%m-%d") if "-" in s else pd.to_datetime(s,format="%d/%m/%Y"))
log.append("Mixed date formats (YYYY-MM-DD and DD/MM/YYYY) standardised")
c["Region"]=c["Region"].str.strip().str.title(); log.append("Region text trimmed and standardised to Title Case")
m=int(c.Region.isna().sum()); c=c.dropna(subset=["Region"]); log.append(f"Rows with missing Region dropped: {m}")
m=int(c.Discount.isna().sum()); c["Discount"]=c["Discount"].fillna(0); log.append(f"Missing Discount filled with 0: {m}")
c["Sales"]=(c.Quantity*c.Unit_Price*(1-c.Discount)).round(2)
c["Cost"]=(c.Quantity*c.Unit_Cost).round(2)
c["Profit"]=(c.Sales-c.Cost).round(2)
c["Month"]=c.Order_Date.dt.strftime("%Y-%m")
c=c.sort_values("Order_Date").reset_index(drop=True)
c=c[["Order_ID","Order_Date","Month","Region","Category","Product","Customer_Segment","Quantity","Unit_Price","Discount","Sales","Cost","Profit"]]
log.append(f"Clean rows: {len(c)}")
c.to_csv(os.path.join(DATA,"cleaned_sales_data.csv"),index=False)
print("\n".join(log))

# ---------- 3. Excel workbook (formulas reference Cleaned_Data) ----------
F="Arial"; hdr_fill=PatternFill("solid",fgColor="1F3864"); hdr_font=Font(name=F,bold=True,color="FFFFFF")
body=Font(name=F); bold=Font(name=F,bold=True)
def head(ws,row,cols):
    for j,t in enumerate(cols,1):
        x=ws.cell(row,j,t); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(horizontal="center",wrap_text=True)
def widths(ws,ws_w):
    for j,v in enumerate(ws_w,1): ws.column_dimensions[get_column_letter(j)].width=v
def fmt(ws,last,ncol,nf):
    for r in range(2,last+1):
        for j in range(1,ncol+1):
            ws.cell(r,j).font=body; ws.cell(r,j).number_format=nf.get(j,"General")
wb=Workbook()
cd=wb.active; cd.title="Cleaned_Data"
head(cd,1,list(c.columns))
for r,row in enumerate(c.itertuples(index=False),2):
    for j,v in enumerate(row,1):
        if j==2: v=v.to_pydatetime()
        x=cd.cell(r,j,v); x.font=body
        if j==2: x.number_format="dd-mmm-yyyy"
        if j in (9,11,12,13): x.number_format="#,##0.00"
        if j==10: x.number_format="0%"
cd.freeze_panes="A2"; cd.auto_filter.ref=f"A1:M{len(c)+1}"; widths(cd,[13,13,10,10,13,15,16,10,11,10,13,13,13])
N=len(c)+1
def rng(col): return f"Cleaned_Data!${col}$2:${col}${N}"

months=sorted(c.Month.unique()); mt=wb.create_sheet("Monthly_Trend")
head(mt,1,["Month","Sales (INR)","Profit (INR)","Orders","Profit Margin","MoM Sales Growth"])
for i,mo in enumerate(months,2):
    mt.cell(i,1,mo)
    mt.cell(i,2,f"=SUMIFS({rng('K')},{rng('C')},A{i})")
    mt.cell(i,3,f"=SUMIFS({rng('M')},{rng('C')},A{i})")
    mt.cell(i,4,f"=COUNTIFS({rng('C')},A{i})")
    mt.cell(i,5,f"=IFERROR(C{i}/B{i},0)")
    mt.cell(i,6,f'=IFERROR(B{i}/B{i-1}-1,"")' if i>2 else '=""')
ML=len(months)+1
fmt(mt,ML,6,{2:"#,##0",3:"#,##0",4:"#,##0",5:"0.0%",6:"0.0%"}); widths(mt,[12,16,16,10,14,18])
ch=LineChart(); ch.title="Monthly Sales Trend"; ch.height=8; ch.width=18
ch.add_data(Reference(mt,min_col=2,min_row=1,max_row=ML),titles_from_data=True); ch.set_categories(Reference(mt,min_col=1,min_row=2,max_row=ML))
mt.add_chart(ch,"H2")

regs=list(c.groupby("Region").Sales.sum().sort_values(ascending=False).index); rp=wb.create_sheet("Region_Performance")
head(rp,1,["Region","Sales (INR)","Profit (INR)","Orders","Profit Margin","Share of Sales"])
RL=len(regs)+1
for i,rg in enumerate(regs,2):
    rp.cell(i,1,rg); rp.cell(i,2,f"=SUMIFS({rng('K')},{rng('D')},A{i})"); rp.cell(i,3,f"=SUMIFS({rng('M')},{rng('D')},A{i})")
    rp.cell(i,4,f"=COUNTIFS({rng('D')},A{i})"); rp.cell(i,5,f"=IFERROR(C{i}/B{i},0)"); rp.cell(i,6,f"=B{i}/SUM($B$2:$B${RL})")
fmt(rp,RL,6,{2:"#,##0",3:"#,##0",4:"#,##0",5:"0.0%",6:"0.0%"}); widths(rp,[12,16,16,10,14,14])
ch=BarChart(); ch.title="Sales by Region"; ch.height=8; ch.width=14
ch.add_data(Reference(rp,min_col=2,min_row=1,max_row=RL),titles_from_data=True); ch.set_categories(Reference(rp,min_col=1,min_row=2,max_row=RL))
rp.add_chart(ch,"H2")

prods=list(c.groupby("Product").Sales.sum().sort_values(ascending=False).index); pp=wb.create_sheet("Product_Performance")
head(pp,1,["Rank","Product","Category","Sales (INR)","Units Sold","Profit (INR)","Profit Margin"])
cat_of=c.drop_duplicates("Product").set_index("Product").Category; PL=len(prods)+1
for i,p in enumerate(prods,2):
    pp.cell(i,1,i-1); pp.cell(i,2,p); pp.cell(i,3,cat_of[p])
    pp.cell(i,4,f"=SUMIFS({rng('K')},{rng('F')},B{i})"); pp.cell(i,5,f"=SUMIFS({rng('H')},{rng('F')},B{i})")
    pp.cell(i,6,f"=SUMIFS({rng('M')},{rng('F')},B{i})"); pp.cell(i,7,f"=IFERROR(F{i}/D{i},0)")
fmt(pp,PL,7,{4:"#,##0",5:"#,##0",6:"#,##0",7:"0.0%"}); widths(pp,[7,16,14,16,12,16,14])
ch=BarChart(); ch.type="bar"; ch.title="Sales by Product"; ch.height=9; ch.width=16
ch.add_data(Reference(pp,min_col=4,min_row=1,max_row=PL),titles_from_data=True); ch.set_categories(Reference(pp,min_col=2,min_row=2,max_row=PL))
pp.add_chart(ch,"I2")

cats=list(c.groupby("Category").Sales.sum().sort_values(ascending=False).index); cp=wb.create_sheet("Category_Performance")
head(cp,1,["Category","Sales (INR)","Profit (INR)","Profit Margin","Share of Sales"]); CL=len(cats)+1
for i,k in enumerate(cats,2):
    cp.cell(i,1,k); cp.cell(i,2,f"=SUMIFS({rng('K')},{rng('E')},A{i})"); cp.cell(i,3,f"=SUMIFS({rng('M')},{rng('E')},A{i})")
    cp.cell(i,4,f"=IFERROR(C{i}/B{i},0)"); cp.cell(i,5,f"=B{i}/SUM($B$2:$B${CL})")
fmt(cp,CL,5,{2:"#,##0",3:"#,##0",4:"0.0%",5:"0.0%"}); widths(cp,[14,16,16,14,14])

ks=wb.create_sheet("KPI_Summary",0)
ks["A1"]="Sales Performance - KPI Summary"; ks["A1"].font=Font(name=F,bold=True,size=14)
ks["A2"]="Synthetic dataset for portfolio purposes. All values are formulas linked to Cleaned_Data."; ks["A2"].font=Font(name=F,italic=True,color="7F7F7F")
head(ks,4,["KPI","Value"])
kp=[("Total Sales (INR)",f"=SUM({rng('K')})","#,##0"),("Total Profit (INR)",f"=SUM({rng('M')})","#,##0"),
    ("Profit Margin","=B6/B5","0.0%"),("Total Orders",f"=COUNTA({rng('A')})","#,##0"),
    ("Units Sold",f"=SUM({rng('H')})","#,##0"),("Average Order Value (INR)","=B5/B8","#,##0"),
    ("Best Month",f"=INDEX(Monthly_Trend!$A$2:$A${ML},MATCH(MAX(Monthly_Trend!$B$2:$B${ML}),Monthly_Trend!$B$2:$B${ML},0))","General"),
    ("Top Region",f"=INDEX(Region_Performance!$A$2:$A${RL},MATCH(MAX(Region_Performance!$B$2:$B${RL}),Region_Performance!$B$2:$B${RL},0))","General"),
    ("Top Product",f"=INDEX(Product_Performance!$B$2:$B${PL},MATCH(MAX(Product_Performance!$D$2:$D${PL}),Product_Performance!$D$2:$D${PL},0))","General")]
for i,(k,f,nf) in enumerate(kp,5):
    ks.cell(i,1,k).font=bold; x=ks.cell(i,2,f); x.font=body; x.number_format=nf; x.alignment=Alignment(horizontal="right")
widths(ks,[30,22])
wb.save(os.path.join(XL,"Sales_Analysis.xlsx"))

# ---------- 4. Stats for README ----------
rs=c.groupby("Region").Sales.sum().sort_values(ascending=False)
ps=c.groupby("Product").Sales.sum().sort_values(ascending=False)
ins=dict(tot=c.Sales.sum(),prof=c.Profit.sum(),orders=len(c),aov=c.Sales.sum()/len(c),
 ms=c.groupby("Month").Sales.sum(),rs=rs,rm=(c.groupby("Region").Profit.sum()/rs).sort_values(ascending=False),
 ps=ps,pu=c.groupby("Product").Quantity.sum().sort_values(ascending=False),
 pm=(c.groupby("Product").Profit.sum()/ps).sort_values(ascending=False),
 cs=c.groupby("Category").Sales.sum().sort_values(ascending=False),
 cm=(c.groupby("Category").Profit.sum()/c.groupby("Category").Sales.sum()).sort_values(ascending=False),log=log)
pickle.dump(ins,open("/tmp/ins.pkl","wb"))
